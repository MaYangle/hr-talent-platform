import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

from talent_hub.model import model_dimensions
from talent_hub.parsing import extract, validate_file


class DocumentTests(unittest.TestCase):
    def test_bad_signature_encoding_and_configured_size_are_rejected(self):
        for name, content, maximum in [
            ("file.pdf", b"not pdf", 100),
            ("file.txt", b"\xff\xfe", 100),
            ("file.txt", b"oversized", 2),
            ("file.exe", b"binary", 100),
        ]:
            with self.subTest(name=name), self.assertRaises(ValueError):
                validate_file(name, content, maximum)

    def test_docx_without_document_part_is_rejected(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("other.xml", "<root/>")
        with self.assertRaises(ValueError):
            validate_file("file.docx", stream.getvalue())

    def test_real_docx_paragraph_and_table_text_is_extracted(self):
        with tempfile.TemporaryDirectory(prefix="talent-docx-") as directory:
            path = Path(directory) / "document.docx"
            doc = Document()
            doc.add_paragraph("alpha")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "beta"
            doc.save(path)
            validate_file(path.name, path.read_bytes())
            pages, ocr = extract(path, path.name)
            self.assertIn("alpha", pages[0]["text"])
            self.assertIn("beta", pages[0]["text"])
            self.assertFalse(ocr)

    def test_image_only_pdf_has_an_explicit_recovery_message(self):
        with tempfile.TemporaryDirectory(prefix="talent-pdf-") as directory:
            path = Path(directory) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.write(path)
            validate_file(path.name, path.read_bytes())
            with self.assertRaisesRegex(ValueError, "OCR"):
                extract(path, path.name, ocr=False)

    def test_shutdown_interrupts_parsing_before_reading_files(self):
        with self.assertRaises(InterruptedError):
            extract("not-read.txt", "not-read.txt", should_stop=lambda: True)


class ModelContractTests(unittest.TestCase):
    def setUp(self):
        self.candidate = {"text": "SQL"}
        self.version = {"criteria": [{"id": "sql", "value": "SQL"}]}
        self.config = {"modelUrl": "http://127.0.0.1:11434/v1", "model": "unit-protocol"}

    def output(self, rows):
        return io.BytesIO(
            json.dumps({"choices": [{"message": {"content": json.dumps(rows)}}]}).encode()
        )

    def call_with(self, rows):
        with (
            patch.dict(os.environ, {}, clear=True),
            patch("talent_hub.model.urllib.request.build_opener") as factory,
        ):
            factory.return_value.open.return_value = self.output(rows)
            return model_dimensions(self.candidate, self.version, self.config)

    def test_external_model_address_is_rejected_before_network_access(self):
        with patch("talent_hub.model.urllib.request.build_opener") as network:
            with self.assertRaises(ValueError):
                model_dimensions(
                    self.candidate,
                    self.version,
                    self.config | {"modelUrl": "https://provider.invalid/v1"},
                )
            network.assert_not_called()

    def test_duplicate_dimensions_and_invalid_unknown_scores_are_rejected(self):
        row = {"criterion_id": "sql", "status": "UNKNOWN", "level": None, "quote": ""}
        for rows in [
            [row, row],
            [row | {"level": 100}],
            [row | {"status": "invalid"}],
            ["not an object"],
        ]:
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                self.call_with(rows)

    def test_valid_response_is_parsed_without_external_service(self):
        row = {"criterion_id": "sql", "status": "MET", "level": 100, "quote": "SQL"}
        self.assertEqual(self.call_with([row]), [row])
