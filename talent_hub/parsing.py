import base64
import html
import io
import json
import os
import subprocess
import tempfile
import threading
import zipfile
from pathlib import Path

_PDFIUM_LOCK = threading.RLock()


def document_preview(data, name):
    ext = Path(name).suffix.lower()
    if ext == ".txt":
        content = "<pre>" + html.escape(data.decode("utf-8-sig")) + "</pre>"
    elif ext == ".docx":
        from docx import Document
        from docx.table import Table

        document = Document(io.BytesIO(data))

        def paragraph_markup(paragraph):
            runs = []
            for item in paragraph.iter_inner_content():
                for run in getattr(item, "runs", [item]):
                    value = html.escape(run.text).replace("\n", "<br>")
                    if run.bold:
                        value = "<strong>" + value + "</strong>"
                    if run.italic:
                        value = "<em>" + value + "</em>"
                    runs.append(value)
                    for picture in run._element.xpath(".//a:blip"):
                        relationship = picture.get(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                        )
                        part = document.part.related_parts.get(relationship)
                        if part and part.content_type in {
                            "image/png",
                            "image/jpeg",
                            "image/gif",
                            "image/webp",
                        }:
                            encoded = base64.b64encode(part.blob).decode("ascii")
                            runs.append(
                                f'<img src="data:{part.content_type};base64,{encoded}" alt="简历内嵌图片">'
                            )
            tag = "h2" if paragraph.style and paragraph.style.name.startswith("Heading") else "p"
            return f"<{tag}>" + ("".join(runs) or "<br>") + f"</{tag}>"

        blocks = []
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                blocks.append(
                    "<table>"
                    + "".join(
                        "<tr>"
                        + "".join(
                            "<td>"
                            + "".join(paragraph_markup(paragraph) for paragraph in cell.paragraphs)
                            + "</td>"
                            for cell in row.cells
                        )
                        + "</tr>"
                        for row in block.rows
                    )
                    + "</table>"
                )
            else:
                blocks.append(paragraph_markup(block))
        content = "".join(blocks)
    else:
        raise ValueError("此格式不支持文字预览")
    return _preview_page(content, name)


def _preview_page(content, name):
    return (
        '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        "<title>" + html.escape(name) + "</title>"
        "<style>body{margin:0;background:#f0f3f5;color:#202b33;font-family:Calibri,'Microsoft YaHei',sans-serif;font-size:16px;line-height:1.85}"
        "article{box-sizing:border-box;max-width:860px;min-height:100vh;margin:auto;padding:40px 48px;background:white;overflow-wrap:anywhere}"
        "p{margin:0 0 14px}h2{font-size:21px;margin:24px 0 12px}table{border-collapse:collapse;width:100%;margin:16px 0}"
        "td{border:1px solid #cbd4da;padding:10px;vertical-align:top}td p:last-child{margin:0}pre{white-space:pre-wrap;font:inherit;margin:0}"
        "img{max-width:100%;height:auto}.pdf-sheet{margin:0 0 24px}.pdf-sheet figcaption{font-size:13px;color:#526673;margin:8px 0}.pdf-sheet img{display:block;width:100%;border:1px solid #d5dde2}"
        "@media(max-width:600px){article{padding:24px 20px}}</style></head><body><article>"
        + content
        + "</article></body></html>"
    ).encode("utf-8")


def pdf_preview(data, image_url, page_number=None):
    with _PDFIUM_LOCK:
        return _pdf_preview(data, image_url, page_number)


def _pdf_preview(data, image_url, page_number):
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(data)
    try:
        count = len(document)
        if not 1 <= count <= 100:
            raise ValueError("PDF 预览支持 1 至 100 页")
        if page_number is None:
            sheets = []
            for index in range(count):
                page = document[index]
                try:
                    width, height = page.get_size()
                finally:
                    page.close()
                sheets.append(
                    f'<figure class="pdf-sheet"><figcaption>第 {index + 1} 页 / 共 {count} 页</figcaption><img src="{html.escape(image_url)}?page={index + 1}" alt="简历第 {index + 1} 页" width="{max(1, round(width))}" height="{max(1, round(height))}" loading="lazy"></figure>'
                )
            return _preview_page("".join(sheets), "简历原文"), "text/html; charset=utf-8"
        if not 1 <= page_number <= count:
            raise ValueError("PDF 页码超出范围")
        page = document[page_number - 1]
        try:
            width, height = page.get_size()
            bitmap = page.render(scale=min(1.8, 1600 / max(1, width), 2400 / max(1, height)))
            try:
                output = io.BytesIO()
                bitmap.to_pil().save(output, format="PNG")
                return output.getvalue(), "image/png"
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        document.close()


def validate_file(name, data, max_bytes=20 * 1024 * 1024):
    ext = Path(name).suffix.lower()
    if not data or len(data) > max_bytes:
        raise ValueError("文件为空或超过上传大小限制")
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("PDF 文件签名无效")
    if ext == ".docx":
        try:
            z = zipfile.ZipFile(io.BytesIO(data))
            if (
                "word/document.xml" not in z.namelist()
                or sum(x.file_size for x in z.infolist()) > 100 * 1024 * 1024
            ):
                raise ValueError()
        except Exception:
            raise ValueError("DOCX 文件损坏或解压体积超限")
    if ext == ".txt":
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("TXT 请另存为 UTF-8 编码")
    if ext not in [".txt", ".pdf", ".docx"]:
        raise ValueError("仅支持 PDF / DOCX / UTF-8 TXT")
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
    }[ext]


def extract(path, name, ocr=True, should_stop=None):
    if should_stop and should_stop():
        raise InterruptedError("服务停止，解析任务重新排队")
    ext = Path(name).suffix.lower()
    used_ocr = False
    if ext == ".txt":
        pages = [{"page": 1, "text": Path(path).read_text(encoding="utf-8-sig")}]
    elif ext == ".docx":
        from docx import Document

        d = Document(path)
        text = "\n".join(
            [p.text for p in d.paragraphs]
            + [" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows]
        )
        pages = [{"page": 1, "text": text}]
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if len(reader.pages) > 100:
            raise ValueError("PDF 超过 100 页，请拆分后重试")
        pages = []
        for n, p in enumerate(reader.pages, 1):
            if should_stop and should_stop():
                raise InterruptedError("服务停止，解析任务重新排队")
            if p.get_contents() and len(p.get_contents().get_data()) > 20 * 1024 * 1024:
                raise ValueError("PDF 页内容解压后过大")
            text = p.extract_text() or ""
            if len(text.strip()) < 12:
                if not ocr:
                    raise ValueError("扫描 PDF 未启用 OCR；可开启后重试或手工建档")
                text = ocr_page(path, n - 1)
                used_ocr = True
            pages.append({"page": n, "text": text})
    else:
        raise ValueError("格式不支持")
    if should_stop and should_stop():
        raise InterruptedError("服务停止，解析任务重新排队")
    if not any(p["text"].strip() for p in pages):
        raise ValueError("未提取到文字，可手工建档")
    return pages, used_ocr


def ocr_page(path, index):
    import pypdfium2 as pdfium

    with tempfile.TemporaryDirectory(prefix="talent-ocr-") as temp:
        image = Path(temp) / "page.png"
        with _PDFIUM_LOCK:
            pdf = pdfium.PdfDocument(path)
            page = pdf[index]
            bitmap = page.render(scale=2)
            bitmap.to_pil().save(image)
            bitmap.close()
            page.close()
            pdf.close()
        if os.name == "nt":
            script = Path(__file__).with_name("ocr.ps1")
            r = subprocess.run(
                [
                    "powershell.exe",
                    "-NoProfile",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    str(script),
                    str(image),
                ],
                capture_output=True,
                timeout=90,
            )
            if r.returncode:
                raise ValueError("Windows 本地 OCR 不可用；请安装系统 OCR 语言包或手工建档")
            return json.loads(r.stdout.decode("utf-8-sig"))["text"]
        r = subprocess.run(
            [
                "tesseract",
                str(image),
                "stdout",
                "-l",
                os.environ.get("TALENT_OCR_LANG", "chi_sim+eng"),
            ],
            capture_output=True,
            timeout=90,
        )
        if r.returncode:
            raise ValueError("本地 Tesseract OCR 失败，请检查语言包")
        return r.stdout.decode("utf-8")
